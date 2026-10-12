"""Score the completed human-verification workbooks (protocol section 5) against the LLM-assisted decisions.

Reads human_check/screening_check_*.xlsx and extraction_check_*.xlsx and sampling_key.json; writes
human_check/results.json and human_check/results.md with, per reviewer and between reviewers:
  title/abstract: n completed, percent agreement, Cohen's kappa with a 95% bootstrap interval, and the sensitivity of
                  the LLM screen against the human decision (the share of human 'include' records the LLM forwarded);
  full text:      the same for eligible (INCLUDE, INCLUDE-O, INCLUDE-A, FIG) versus EXCLUDE, plus exact-category kappa;
  extraction:     share of ID and OOD values that match the extracted value (same number, allowing percent/fraction).
Disagreements are listed for resolution by discussion; record the resolved values with scripts/record.py.
Usage (from paper/review):  python3 scripts/human_check_stats.py"""
import csv
import glob
import json
import os
import random

from openpyxl import load_workbook

D = "human_check"
key = json.load(open(f"{D}/sampling_key.json"))

ta_inc = {l.split()[0] for l in open("screening/ta_include.txt") if l.strip()}
if os.path.exists("screening/chase_ta_include.txt"):
    ta_inc |= {l.split()[0] for l in open("screening/chase_ta_include.txt") if l.strip()}
llm_ta = {r: ("include" if r in ta_inc else "exclude") for r in key["ta_sample"]}
llm_ft = {r["rid"]: r["decision"] for r in csv.DictReader(open("extraction/ft_decisions.csv"))}
comps = {f"{r['rid']}#{r['comp']}": r for r in csv.DictReader(open("extraction/comparisons.csv"))}
ELIG = {"INCLUDE", "INCLUDE-O", "INCLUDE-A", "FIG"}


def kappa(a, b):
    """Cohen's kappa for two equal-length label lists."""
    n = len(a)
    if n == 0:
        return None
    cats = sorted(set(a) | set(b))
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in cats)
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def kappa_ci(a, b, draws=2000, seed=1):
    rng = random.Random(seed)
    n = len(a)
    if n < 10:
        return None, None
    ks = []
    for _ in range(draws):
        idx = [rng.randrange(n) for _ in range(n)]
        k = kappa([a[i] for i in idx], [b[i] for i in idx])
        if k is not None:
            ks.append(k)
    ks.sort()
    return ks[int(0.025 * len(ks))], ks[int(0.975 * len(ks)) - 1]


def summary(pairs, positive):
    a, b = [x for x, _ in pairs], [y for _, y in pairs]
    k = kappa(a, b)
    lo, hi = kappa_ci(a, b)
    human_pos = [i for i, x in enumerate(a) if x == positive]
    sens = (sum(b[i] == positive for i in human_pos) / len(human_pos)) if human_pos else None
    return dict(n=len(pairs), agreement=(sum(x == y for x, y in pairs) / len(pairs)) if pairs else None,
                kappa=k, kappa_ci=[lo, hi], human_positive=len(human_pos), llm_sensitivity=sens)


def read_sheet(path, sheet, id_col, cols):
    ws = load_workbook(path, data_only=True)[sheet]
    hdr = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(hdr, row))
        vals = {c: rec.get(c) for c in cols}
        if any(v not in (None, "") for v in vals.values()):
            out[str(rec[id_col])] = vals
    return out


res = {"screening": {}, "full_text": {}, "extraction": {}, "between_reviewers": {}}
lines = ["# Human verification results", ""]
ta_by, ft_by = {}, {}
for path in sorted(glob.glob(f"{D}/screening_check_*.xlsx")):
    who = os.path.basename(path)[len("screening_check_"):-5]
    ta = read_sheet(path, "TA sample", "Record", ["Your decision"])
    ft = read_sheet(path, "FT check", "Record", ["Your decision", "Reason (if excluded)"])
    ta_by[who] = {r: str(v["Your decision"]).strip().lower() for r, v in ta.items() if v["Your decision"]}
    ft_by[who] = {r: str(v["Your decision"]).strip().upper() for r, v in ft.items() if v["Your decision"]}
    s_ta = summary([(ta_by[who][r], llm_ta[r]) for r in ta_by[who] if r in llm_ta], "include")
    pairs_ft = [(ta, llm_ft[r]) for r, ta in ft_by[who].items() if r in llm_ft]
    s_ft = summary([("eligible" if x in ELIG else "excluded", "eligible" if y in ELIG else "excluded") for x, y in pairs_ft], "eligible")
    s_ft["kappa_exact_category"] = kappa([x for x, _ in pairs_ft], [y for _, y in pairs_ft])
    res["screening"][who], res["full_text"][who] = s_ta, s_ft
    lines += [f"## Reviewer: {who}", "",
              f"- Title/abstract: {s_ta['n']} of {len(key['ta_sample'])} completed; agreement with the LLM screen "
              f"{s_ta['agreement']}; Cohen's kappa {s_ta['kappa']} (95% CI {s_ta['kappa_ci']}); "
              f"LLM sensitivity for records this reviewer included {s_ta['llm_sensitivity']} (n = {s_ta['human_positive']}).",
              f"- Full text: {s_ft['n']} of {len(key['ft_records'])} completed; eligible-vs-excluded kappa {s_ft['kappa']} "
              f"(95% CI {s_ft['kappa_ci']}); exact-category kappa {s_ft['kappa_exact_category']}.", ""]
    dis = [r for r in ta_by[who] if r in llm_ta and ta_by[who][r] != llm_ta[r]]
    lines += [f"  Title/abstract disagreements to discuss ({len(dis)}): " + ", ".join(dis[:300]), ""]
    dis = [r for r in ft_by[who] if r in llm_ft and (ft_by[who][r] in ELIG) != (llm_ft[r] in ELIG)]
    lines += [f"  Full-text eligibility disagreements to discuss ({len(dis)}): " + ", ".join(dis), ""]

who = sorted(ta_by)
if len(who) >= 2:
    a, b = who[0], who[1]
    common = [r for r in ta_by[a] if r in ta_by[b]]
    res["between_reviewers"]["title_abstract"] = summary([(ta_by[a][r], ta_by[b][r]) for r in common], "include")
    common = [r for r in ft_by[a] if r in ft_by[b]]
    res["between_reviewers"]["full_text"] = summary([("eligible" if ft_by[a][r] in ELIG else "excluded",
                                                      "eligible" if ft_by[b][r] in ELIG else "excluded") for r in common], "eligible")
    lines += ["## Between reviewers", "", f"- Title/abstract: {res['between_reviewers']['title_abstract']}",
              f"- Full text: {res['between_reviewers']['full_text']}", ""]


def num(x):
    try:
        return float(str(x).replace("%", "").replace(",", ".").strip())
    except (TypeError, ValueError):
        return None


def same(h, e):
    if h is None or e is None:
        return None
    return any(abs(h * f - e) <= 0.005 * max(1.0, abs(e)) for f in (1, 100, 0.01))


for path in sorted(glob.glob(f"{D}/extraction_check_*.xlsx")):
    who_e = os.path.basename(path)[len("extraction_check_"):-5]
    ws = load_workbook(path, data_only=True)["Extraction"]
    hdr = [c.value for c in ws[1]]
    rows, diffs = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(hdr, row))
        k = f"{rec['Record']}#{rec['Comparison no.']}"
        if k not in comps or (rec.get("ID value") in (None, "") and rec.get("OOD value") in (None, "")):
            continue
        c = comps[k]
        m_id, m_ood = same(num(rec.get("ID value")), num(c["id"])), same(num(rec.get("OOD value")), num(c["ood"]))
        rows.append((m_id, m_ood))
        if m_id is False or m_ood is False:
            diffs.append(f"{k}: human ID {rec.get('ID value')} / OOD {rec.get('OOD value')} vs extracted {c['id']} / {c['ood']}")
    n = len(rows)
    ok = sum(1 for a, b in rows if a is not False and b is not False)
    res["extraction"][who_e] = dict(n=n, of=len(key["comparisons"]), both_values_match=(ok / n) if n else None, discrepancies=diffs)
    lines += [f"## Extraction check: {who_e}", "", f"- {n} of {len(key['comparisons'])} rows completed; both values match "
              f"in {ok} ({(100 * ok / n) if n else 0:.0f}%).", "", "Discrepancies to resolve against the source:", ""]
    lines += [f"- {d}" for d in diffs] + [""]

json.dump(res, open(f"{D}/results.json", "w"), indent=1)
open(f"{D}/results.md", "w").write("\n".join(lines) + "\n")
print("\n".join(lines))
