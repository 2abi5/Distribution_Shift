"""Build the human-verification workbooks required by protocol section 5 (one copy per reviewer).

  screening_check_<reviewer>.xlsx
      'TA sample' - a random 20% of all title/abstract decisions (database search and citation chasing), in random
                    order, so that any completed prefix is itself a random sample;
      'FT check'  - every full-text decision (protocol: 100%), in random order.
  extraction_check_<reviewer>.xlsx
      'Extraction' - a random 20% of extracted ID-OOD comparisons; the reviewer reads the ID and OOD values from the
                     cited table or figure. Extracted values are NOT shown (blind double extraction).

The LLM-assisted decisions and values are deliberately absent from the workbooks. human_check_stats.py joins them
after the reviewers have finished and computes Cohen's kappa, agreement and extraction error rates.
Usage (from paper/review):  python3 scripts/make_human_check.py [--seed 20261002] [--reviewers "Abinash Pant,Binayak Prakash Mishra"]"""
import csv
import json
import os
import random
import re
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SEED = int(sys.argv[sys.argv.index("--seed") + 1]) if "--seed" in sys.argv else 20261002
REVIEWERS = (sys.argv[sys.argv.index("--reviewers") + 1] if "--reviewers" in sys.argv
             else "Abinash Pant,Binayak Prakash Mishra").split(",")
OUT = "human_check"
os.makedirs(OUT, exist_ok=True)

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F4E78")
INPUT = PatternFill("solid", fgColor="FFF2CC")        # cells the reviewer fills in
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")


def clean(s):
    s = re.sub(r"<[^>]+>", "", s or "")
    s = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", s)        # characters Excel rejects
    return re.sub(r"\s+", " ", s).strip()


# ------------------------------------------------------------------ data
pool = {}
for line in [*open("search/pool_v3.jsonl"), *(open("search/pool_extra.jsonl") if os.path.exists("search/pool_extra.jsonl") else [])]:
    r = json.loads(line)
    pool[r["rid"]] = r
chase_ids = set(json.load(open("search/chase_screen_ids.json"))) if os.path.exists("search/chase_screen_ids.json") else set()
chased = {}
if chase_ids:
    for line in open("search/chase_2026-10-02.jsonl"):
        r = json.loads(line)
        if r["wid"] in chase_ids:
            chased[r["wid"]] = r

ta_records = []
for rid, p in pool.items():
    if rid.startswith("R"):
        ta_records.append(dict(id=rid, source="database search", title=clean(p.get("title")), abstract=clean(p.get("abstract")),
                               year=p.get("year", ""), venue=clean(p.get("venue")), doi=p.get("doi") or ""))
for wid, p in chased.items():
    ta_records.append(dict(id=wid, source="citation chasing", title=clean(p.get("title")), abstract=clean(p.get("abstract")),
                           year=p.get("year", ""), venue=clean(p.get("venue")), doi=p.get("doi") or ""))
rng = random.Random(SEED)
ta_records.sort(key=lambda r: r["id"])
ta_sample = rng.sample(ta_records, round(0.2 * len(ta_records)))      # random order = sampling order

ft = list(csv.DictReader(open("extraction/ft_decisions.csv")))
chase_meta = {}
if os.path.exists("search/chase_2026-10-02.jsonl"):
    for line in open("search/chase_2026-10-02.jsonl"):
        r = json.loads(line)
        chase_meta[r["wid"]] = r
idmap = {}
if os.path.exists("search/chase_ids.csv"):
    idmap = {r["rid"]: r["wid"] for r in csv.DictReader(open("search/chase_ids.csv"))}


def meta(rid):
    p = pool.get(rid) or chase_meta.get(idmap.get(rid, rid)) or {}
    return clean(p.get("title")), p.get("doi") or "", p.get("year", "")


ft_rows = []
for d in ft:
    title, doi, year = meta(d["rid"])
    txt = f"paper/review/fulltext/txt/{d['rid']}.txt"
    ft_rows.append(dict(id=d["rid"], title=title, year=year, doi=doi,
                        local=txt if os.path.exists(txt.replace("paper/review/", "")) else "(not retrieved: use the DOI)"))
rng.shuffle(ft_rows)

comps = list(csv.DictReader(open("extraction/comparisons.csv")))
comp_sample = rng.sample(comps, round(0.2 * len(comps)))


# ------------------------------------------------------------------ sheet helpers
def header(ws, cols, widths):
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font, cell.fill, cell.alignment = HEAD, HFILL, WRAP
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32


def put(ws, i, values, inputs, height=None):
    for j, v in enumerate(values, 1):
        cell = ws.cell(row=i, column=j, value=v)
        cell.font, cell.alignment, cell.border = FONT, WRAP, THIN
        if j in inputs:
            cell.fill = INPUT
    if height:
        ws.row_dimensions[i].height = height


def readme(ws, title, lines, example_cols=None, example=None):
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    r = 3
    for ln in lines:
        c = ws.cell(row=r, column=1, value=ln)
        c.font = BOLD if ln.endswith(":") else FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    if example:
        r += 1
        ws.cell(row=r, column=1, value="Example of a completed row (format only; not part of the sample):").font = BOLD
        r += 1
        for j, (h, v) in enumerate(zip(example_cols, example), 1):
            h_cell = ws.cell(row=r, column=j, value=h)
            h_cell.font, h_cell.fill, h_cell.alignment = HEAD, HFILL, WRAP
            v_cell = ws.cell(row=r + 1, column=j, value=v)
            v_cell.font, v_cell.alignment = FONT, WRAP
    ws.column_dimensions["A"].width = 110
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 18


CRITERIA = [
    "Eligibility (protocol section 4), in short:",
    "Include if ALL hold: I1 primary study training a supervised ML model for a crop-production task (disease/pest/abiotic "
    "stress; weed-crop discrimination; detection/counting/segmentation of plants or organs; phenotypic traits; crop-type or "
    "cropland mapping; yield); I2 performance reported on an in-distribution (ID) test set AND on at least one "
    "out-of-distribution (OOD) test set from a different acquisition setting, location, time, sensor/platform or biological "
    "material, with the same metric; I3 the unadapted (source-only) model's OOD result is reported.",
    "Exclude: E1 livestock, aquaculture, forestry, post-harvest, soil-only or generic land cover; E2 shift only synthetic "
    "(noise, simulation, synthetic-to-real); E3 random/k-fold splits only; E4 review, editorial, non-English, or full text "
    "unavailable and abstract lacks ID and OOD values; E5 OOD label space disjoint from training.",
    "Title/abstract stage is inclusive: mark 'include' when the record plausibly reports an evaluation on a domain other "
    "than the training domain, or names domain adaptation, domain generalisation, cross-domain testing or transferability.",
]


# ------------------------------------------------------------------ screening workbook
def screening_book(reviewer):
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    readme(ws, f"Screening check - reviewer: {reviewer}", [
        "Purpose:",
        "Protocol section 5 requires an independent human check of the LLM-assisted screening. Work on your own; do not "
        "discuss records with the other reviewer until both files are complete. The LLM decisions are not in this file.",
        "What to fill in (yellow cells only):",
        "Sheet 'TA sample': column 'Your decision' (include / exclude) for each title and abstract; 'Notes' is optional.",
        "Sheet 'FT check': open the full text (local path or DOI), then 'Your decision' and, for exclusions, 'Reason'.",
        "Order:",
        f"Rows are in random order. If you cannot finish all {len(ta_sample)} title/abstract rows, complete them from the "
        "top: any completed top portion is itself a random sample. Report how many you completed.",
        "Do not change any other column, sort the sheet or delete rows.",
        *CRITERIA,
        "When both reviewers are done: save the files in paper/review/human_check/ and run",
        "    python3 scripts/human_check_stats.py",
        "which computes Cohen's kappa (reviewer vs LLM, reviewer vs reviewer), agreement, and the list of disagreements to "
        "resolve by discussion.",
    ], ["Order", "Record", "Source", "Title", "Abstract", "Year", "Your decision", "Notes"],
        [1, "R01234", "database search", "Unsupervised domain adaptation for crop mapping in new regions",
         "We train a model in France and test it in Spain ...", 2024, "include", "OOD = new country"])

    ta = wb.create_sheet("TA sample")
    cols = ["Order", "Record", "Source", "Title", "Abstract", "Year", "Venue", "DOI", "Your decision", "Notes"]
    header(ta, cols, [7, 13, 13, 40, 90, 6, 22, 22, 13, 25])
    dv = DataValidation(type="list", formula1='"include,exclude"', allow_blank=True)
    ta.add_data_validation(dv)
    for i, r in enumerate(ta_sample, 2):
        put(ta, i, [i - 1, r["id"], r["source"], r["title"], r["abstract"][:30000] or "(no abstract)", r["year"], r["venue"],
                    r["doi"], None, None], inputs={9, 10}, height=120)
    dv.add(f"I2:I{len(ta_sample) + 1}")
    ta.auto_filter.ref = f"A1:J{len(ta_sample) + 1}"

    fs = wb.create_sheet("FT check")
    cols = ["Order", "Record", "Title", "Year", "DOI", "Full text on this computer", "Your decision", "Reason (if excluded)", "Notes"]
    header(fs, cols, [7, 13, 50, 6, 26, 34, 14, 14, 30])
    dv2 = DataValidation(type="list", formula1='"INCLUDE,INCLUDE-O,INCLUDE-A,FIG,EXCLUDE"', allow_blank=True)
    dv3 = DataValidation(type="list", formula1='"I1,I2,I3,E1,E2,E3,E4,E5,NR"', allow_blank=True)
    fs.add_data_validation(dv2)
    fs.add_data_validation(dv3)
    for i, r in enumerate(ft_rows, 2):
        put(fs, i, [i - 1, r["id"], r["title"], r["year"], r["doi"], r["local"], None, None, None], inputs={7, 8, 9}, height=30)
    dv2.add(f"G2:G{len(ft_rows) + 1}")
    dv3.add(f"H2:H{len(ft_rows) + 1}")
    fs["G1"].comment = Comment("INCLUDE: full text, ID and OOD of the same source-only model. INCLUDE-O: source-only "
                               "and target-trained results on the OOD test set but no ID result. INCLUDE-A: no full "
                               "text, but the abstract gives paired ID and OOD values. FIG: eligible, values only in "
                               "figures. EXCLUDE: give the reason.", "protocol")
    fs.auto_filter.ref = f"A1:I{len(ft_rows) + 1}"
    path = f"{OUT}/screening_check_{reviewer.split()[0].lower()}.xlsx"
    wb.save(path)
    return path


# ------------------------------------------------------------------ extraction workbook
def extraction_book(reviewer):
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    readme(ws, f"Extraction check - reviewer: {reviewer}", [
        "Purpose:",
        "Protocol section 5: double extraction of a random 20% of comparisons. For each row, open the study, go to the "
        "table or figure named in 'Where to look', and copy the ID and OOD values exactly as printed (with the units the "
        "paper uses, e.g. 91.2 or 0.912). The values extracted by the LLM-assisted reviewer are not in this file.",
        "What to fill in (yellow cells only):",
        "'ID value' and 'OOD value' as printed; 'Adapted value' only when the row describes a mitigation; 'Notes' for "
        "anything that does not fit (value not found, different model, figure only, etc.).",
        "If a value is only in a figure, read it as well as you can and write 'figure' in Notes.",
        "Do not change any other column. Rows are in random order; complete them from the top if time is short.",
        "Then run:  python3 scripts/human_check_stats.py",
    ], ["Order", "Record", "Study", "Comparison", "Model", "Metric", "Where to look", "ID value", "OOD value", "Notes"],
        [1, "R00004", "Lottes et al. 2018", "time: Bonn 2016 -> Bonn 2017", "Vanilla FCN", "avg F1 object-wise (%)",
         "Table V", 96.7, 72.5, ""])
    sh = wb.create_sheet("Extraction")
    cols = ["Order", "Record", "Comparison no.", "Study", "DOI", "Task", "Shift", "Shift description", "Model", "Metric",
            "Mitigation", "Where to look", "ID value", "OOD value", "Adapted value", "Notes"]
    header(sh, cols, [7, 11, 9, 20, 24, 6, 11, 38, 26, 18, 12, 18, 10, 10, 10, 28])
    for i, c in enumerate(comp_sample, 2):
        title, doi, _ = meta(c["rid"])
        mit = "" if c["mit"] in ("", "none") else f"{c['mit']} {c['mit_name']}".strip()
        put(sh, i, [i - 1, c["rid"], int(c["comp"]), c["study"], doi, c["task"], c["shift"], c["shift_desc"], c["model"],
                    c["metric"], mit, c["loc"], None, None, None, None], inputs={13, 14, 15, 16}, height=30)
    sh.auto_filter.ref = f"A1:P{len(comp_sample) + 1}"
    path = f"{OUT}/extraction_check_{reviewer.split()[0].lower()}.xlsx"
    wb.save(path)
    return path


key = {"seed": SEED, "ta_sample": [r["id"] for r in ta_sample], "ft_records": [r["id"] for r in ft_rows],
       "comparisons": [f"{c['rid']}#{c['comp']}" for c in comp_sample],
       "n_ta_population": len(ta_records), "n_ft": len(ft_rows), "n_comparisons": len(comps)}
json.dump(key, open(f"{OUT}/sampling_key.json", "w"), indent=1)
for rv in REVIEWERS:
    print(screening_book(rv.strip()), extraction_book(rv.strip()))
print(f"TA sample {len(ta_sample)} of {len(ta_records)}; FT records {len(ft_rows)}; comparisons {len(comp_sample)} of {len(comps)}")
